from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document


SOURCE = Path(
    "/Users/glebkorotkov/github/lifechange/lifeisgame/работа над курсовой/"
    "Курсовая работа Глеба Короткова — диаграммы в тексте.docx"
)
OUTPUT = Path(
    "/Users/glebkorotkov/github/lifechange/lifeisgame/работа над курсовой/"
    "Курсовая работа Глеба Короткова — правки UI и форматирование.docx"
)


UIKIT_TEXT = (
    "UIKit выбран основным UI–фреймворком благодаря императивной модели управления интерфейсом. "
    "Она позволяет явно контролировать создание и конфигурацию UIView, жизненный цикл "
    "UIViewController, обработку событий, иерархию представлений и навигационные переходы. "
    "Такой уровень управления важен для форм с валидацией, календаря со сложной компоновкой, "
    "многошагового создания задач и координаторной навигации Luma. Программная вёрстка в UIKit "
    "также даёт возможность точно задавать ограничения Auto Layout, порядок обновления элементов "
    "и реакцию экрана на действия пользователя. Поэтому UIKit используется как единая основа "
    "основной части интерфейса."
)

SWIFTUI_TEXT = (
    "SwiftUI выбран для экрана достижений и интерфейса Live Activity, поскольку эти части "
    "преимущественно отображают состояние и требуют автоматического обновления представления при "
    "изменении данных. Декларативная модель SwiftUI позволяет описывать соответствие между "
    "состоянием и интерфейсом без ручного управления каждым обновлением элементов [34]. Для "
    "экрана достижений это упрощает отображение списка, прогресса и состояний блокировки, а для "
    "Live Activity обеспечивает предусмотренную ActivityKit и WidgetKit модель представления на "
    "экране блокировки и в Dynamic Island [33; 41]. Экран достижений встраивается в основной "
    "UIKit-интерфейс через UIHostingController. Таким образом, UIKit отвечает за сценарии со "
    "сложным управлением экраном и навигацией, а SwiftUI применяется в изолированных "
    "представлениях, содержание которых определяется текущим состоянием данных."
)


TOC_PAGE_MAP = {
    "ВВЕДЕНИЕ": 3,
    "1 Анализ предметной области": 5,
    "1.1 Особенности планирования задач с учётом состояния пользователя": 5,
    "1.2 Анализ приложений-аналогов": 6,
    "1.3 Нефункциональные требования": 8,
    "1.4 Функциональные возможности приложения": 10,
    "1.5 Модель предметной области": 17,
    "2 Проектирование": 19,
    "2.1 Выбор платформы и средств разработки": 19,
    "2.2 Архитектурный подход": 20,
    "2.3 Хранение данных и системные интеграции": 22,
    "2.4 Проектирование пользовательского интерфейса": 24,
    "3 Реализация и проверка приложения": 25,
    "3.1 Учётная запись, профиль и сессия": 25,
    "3.2 Календарь и управление элементами": 28,
    "3.3 Автоматическое планирование и граница 23:00": 30,
    "3.4 Уведомления и дневник": 34,
    "3.5 Фокус-сессия и Live Activity": 36,
    "3.6 Анализ дня, статистика и достижения": 37,
    "3.7 Пользовательское исследование восприятия интерфейса": 42,
    "ЗАКЛЮЧЕНИЕ": 46,
    "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ И ЛИТЕРАТУРЫ": 47,
    "ПРИЛОЖЕНИЕ А": 52,
    "ПРИЛОЖЕНИЕ Б": 55,
}


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def rewrite_text(doc: Document) -> None:
    replaced_uikit = False
    replaced_swiftui = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("UIKit выбран основным UI"):
            replace_paragraph(paragraph, UIKIT_TEXT)
            replaced_uikit = True
        elif text.startswith("SwiftUI применён точечно"):
            replace_paragraph(paragraph, SWIFTUI_TEXT)
            replaced_swiftui = True
    if not replaced_uikit or not replaced_swiftui:
        raise RuntimeError(
            f"Expected UI paragraphs were not found: UIKit={replaced_uikit}, SwiftUI={replaced_swiftui}"
        )


def update_static_toc(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if "\t" not in text:
            continue
        title, _, old_page = text.rpartition("\t")
        if title in TOC_PAGE_MAP and old_page.isdigit():
            replace_paragraph(paragraph, f"{title}\t{TOC_PAGE_MAP[title]}")


def patch_ooxml(input_path: Path, output_path: Path) -> None:
    temp_dir = Path(tempfile.mkdtemp(prefix="luma-coursework-format-"))
    try:
        with zipfile.ZipFile(input_path) as source:
            source.extractall(temp_dir)

        word_dir = temp_dir / "word"
        for xml_path in word_dir.rglob("*.xml"):
            data = xml_path.read_text(encoding="utf-8")
            data = data.replace("—", "–")
            data = re.sub(r'(<w:sz(?:Cs)?\b[^>]*\bw:val=")[^"]+("[^>]*/?>)', r"\g<1>24\2", data)
            xml_path.write_text(data, encoding="utf-8")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for path in sorted(temp_dir.rglob("*")):
                if path.is_file():
                    target.write(path, path.relative_to(temp_dir))
    finally:
        shutil.rmtree(temp_dir)


def main() -> None:
    temp_docx = Path(tempfile.mkstemp(prefix="luma-coursework-edit-", suffix=".docx")[1])
    try:
        doc = Document(SOURCE)
        rewrite_text(doc)
        update_static_toc(doc)
        doc.save(temp_docx)
        patch_ooxml(temp_docx, OUTPUT)
    finally:
        temp_docx.unlink(missing_ok=True)
    print(OUTPUT)


if __name__ == "__main__":
    main()
