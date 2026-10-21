from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt


SOURCE = Path(
    "/Users/glebkorotkov/github/lifechange/lifeisgame/работа над курсовой/"
    "Курсовая работа Глеба Короткова — нефункциональные требования текстом.docx"
)
OUTPUT = Path(
    "/Users/glebkorotkov/github/lifechange/lifeisgame/работа над курсовой/"
    "Курсовая работа Глеба Короткова — функциональные возможности.docx"
)
IMAGE_DIR = Path(
    "/Users/glebkorotkov/github/lifechange/lifeisgame/.codex-work/luma-use-cases"
)


def paragraph_after(paragraph: Paragraph, text: str = "", style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    result = Paragraph(element, paragraph._parent)
    result.style = style
    if text:
        result.add_run(text)
    return result


def format_body(paragraph: Paragraph, *, keep_with_next: bool = False) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.widow_control = True


def add_body(anchor: Paragraph, text: str, *, keep_with_next: bool = False) -> Paragraph:
    paragraph = paragraph_after(anchor, text)
    format_body(paragraph, keep_with_next=keep_with_next)
    return paragraph


def add_figure(anchor: Paragraph, image_path: Path, caption: str) -> Paragraph:
    image_paragraph = paragraph_after(anchor)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = Pt(0)
    image_paragraph.paragraph_format.space_before = Pt(4)
    image_paragraph.paragraph_format.space_after = Pt(0)
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.35))

    caption_paragraph = paragraph_after(image_paragraph, caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = Pt(0)
    caption_paragraph.paragraph_format.space_before = Pt(0)
    caption_paragraph.paragraph_format.space_after = Pt(6)
    caption_paragraph.paragraph_format.widow_control = True
    return caption_paragraph


def replace_text_nodes(document: Document, old: str, new: str) -> None:
    for node in document._element.xpath(".//w:t"):
        if node.text and old in node.text:
            node.text = node.text.replace(old, new)


def main() -> None:
    document = Document(SOURCE)
    paragraphs = document.paragraphs

    section_heading = next(
        paragraph
        for paragraph in paragraphs
        if paragraph.text.strip() == "1.4 Модель вариантов использования"
    )
    next_heading = next(
        paragraph
        for paragraph in paragraphs
        if paragraph.text.strip() == "1.5 Модель предметной области"
    )

    # Remove the previous body of section 1.4, including its raster figures.
    current = section_heading._p.getnext()
    while current is not None and current is not next_heading._p:
        following = current.getnext()
        current.getparent().remove(current)
        current = following

    # Update both the visible heading and the cached TOC entry.
    replace_text_nodes(
        document,
        "1.4 Модель вариантов использования",
        "1.4 Функциональные возможности приложения",
    )

    anchor = section_heading
    anchor = add_body(
        anchor,
        "Функциональные требования к приложению определены техническим заданием. "
        "В курсовой работе они представлены диаграммами вариантов использования: "
        "такой формат связывает возможности приложения с целями пользователя и "
        "показывает участие системных служб iOS без повторения требований отдельным перечнем.",
    )
    anchor = add_body(
        anchor,
        "Основным актором является пользователь. До входа он может зарегистрироваться "
        "или авторизоваться, а после входа получает доступ к планированию, дневнику, "
        "фокус-сессиям и анализу состояния. Системные службы iOS показаны как "
        "вспомогательные акторы только в тех случаях, когда они действительно участвуют "
        "во взаимодействии.",
    )

    anchor = add_body(
        anchor,
        "Начало работы включает регистрацию и несколько способов входа в учётную запись "
        "(рисунок 1.1).",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-1.png",
        "Рисунок 1.1 - Регистрация и авторизация",
    )
    anchor = add_body(
        anchor,
        "Новый пользователь создаёт учётную запись, указывая данные, предусмотренные "
        "формой регистрации. Для последующего входа доступны электронная почта и пароль. "
        "После первичной авторизации пользователь может включить вход через Face ID; "
        "системная служба выполняет биометрическую проверку, а приложение сохраняет "
        "контроль над состоянием пользовательской сессии.",
    )

    anchor = add_body(
        anchor,
        "После авторизации пользователь может просмотреть и изменить параметры профиля "
        "(рисунок 1.2).",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-2.png",
        "Рисунок 1.2 - Профиль пользователя",
    )
    anchor = add_body(
        anchor,
        "В профиле отображаются сведения об учётной записи, уровень и накопленный опыт. "
        "Пользователь может изменить аватар, настроить биометрический вход или выйти из "
        "учётной записи. Служба выбора изображений предоставляет выбранный файл, а Face ID "
        "используется только для подтверждения личности.",
    )

    anchor = add_body(
        anchor,
        "Работа с напоминаниями и журналом уведомлений представлена на рисунке 1.3.",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-3.png",
        "Рисунок 1.3 - Настройка и получение уведомлений",
    )
    anchor = add_body(
        anchor,
        "Пользователь самостоятельно включает или отключает уведомления и получает "
        "напоминания о запланированных действиях. Системная служба iOS отвечает за доставку "
        "уведомления. История уже доставленных сообщений хранится и просматривается внутри "
        "приложения, поэтому этот сценарий не зависит от системного центра уведомлений.",
    )

    anchor = add_body(
        anchor,
        "Календарь объединяет собственные задачи пользователя и внешние события, полученные "
        "из системного календаря (рисунок 1.4).",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-4.png",
        "Рисунок 1.4 - Просмотр календаря и управление задачами",
    )
    anchor = add_body(
        anchor,
        "После выдачи разрешения приложение читает события системного календаря и учитывает "
        "занятые интервалы при составлении плана. Пользователь выбирает день, просматривает "
        "его расписание, применяет фильтр и управляет собственными задачами. Импортированные "
        "события доступны только для чтения и не редактируются в приложении.",
    )

    anchor = add_body(
        anchor,
        "Создание задач и событий связано с автоматическим подбором времени, что отражено "
        "на рисунке 1.5.",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-5.png",
        "Рисунок 1.5 - Создание задач и автоматическое планирование",
    )
    anchor = add_body(
        anchor,
        "Событие создаётся для заранее выбранного времени. Для обычной задачи пользователь "
        "задаёт продолжительность и ограничения, после чего приложение ищет подходящий "
        "свободный интервал. Сложная задача дополнительно разбивается на подзадачи, которые "
        "могут быть распределены по нескольким доступным периодам.",
    )
    anchor = add_body(
        anchor,
        "Центральный сценарий автоматического планирования начинается после заполнения "
        "параметров задачи. Приложение сопоставляет её длительность с уже занятыми и "
        "свободными интервалами, выбирает допустимое время и добавляет результат в план дня. "
        "Если непрерывного интервала недостаточно, для сложной задачи используются отдельные "
        "подзадачи; созданное пользователем событие с фиксированным временем при этом не "
        "переносится.",
    )

    anchor = add_body(
        anchor,
        "Функции фиксации состояния и отображения прогресса собраны на рисунке 1.6.",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-6.png",
        "Рисунок 1.6 - Достижения и дневник состояния",
    )
    anchor = add_body(
        anchor,
        "В дневнике пользователь отмечает эмоциональное состояние и сведения о сне. "
        "Эти записи используются при последующем анализе дня. Раздел достижений показывает "
        "текущий уровень, накопленный опыт и условия уже открытых достижений. Два набора "
        "сценариев логически независимы, но связаны с одной учётной записью.",
    )

    anchor = add_body(
        anchor,
        "Фокус-сессия помогает выполнять выбранную задачу без постоянного возврата в "
        "основной интерфейс приложения (рисунок 1.7).",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-7.png",
        "Рисунок 1.7 - Фокус-сессия и Live Activity",
    )
    anchor = add_body(
        anchor,
        "Пользователь запускает сессию для выбранной задачи и наблюдает её ход в Live "
        "Activity. Через системное представление можно завершить задачу, после чего действие "
        "передаётся приложению. Служба Live Activities отображает текущее состояние и "
        "принимает команду, но инициатором сценария и получателем результата остаётся "
        "пользователь.",
    )

    anchor = add_body(
        anchor,
        "Итоговый анализ дня связывает данные календаря с выполненными задачами и записями "
        "о состоянии пользователя (рисунок 1.8).",
        keep_with_next=True,
    )
    anchor = add_figure(
        anchor,
        IMAGE_DIR / "use-case-8.png",
        "Рисунок 1.8 - Анализ состояния пользователя",
    )
    anchor = add_body(
        anchor,
        "Приложение рассчитывает показатели нагрузки, сопоставляет их с продолжительностью "
        "сна и эмоциональными отметками, а затем формирует рекомендации. Системный календарь "
        "передаёт внешние события, необходимые для оценки занятого и свободного времени. "
        "Пользователь получает результат в виде сводки и может перейти к следующей "
        "фокус-сессии или скорректировать план.",
    )
    add_body(
        anchor,
        "Таким образом, диаграммы описывают путь от первого входа до ежедневного "
        "планирования, выполнения задач и оценки результатов. Далее эти сценарии используются "
        "для выделения сущностей предметной области.",
    )

    replace_text_nodes(document, "рисунке 1.11", "рисунке 1.9")
    replace_text_nodes(
        document,
        "Рисунок 1.11 - Модель предметной области",
        "Рисунок 1.9 - Модель предметной области",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
