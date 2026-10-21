#!/usr/bin/env python3
"""Extract compact, reusable review artifacts from DOCX files."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def paragraph_record(paragraph, index: int) -> dict:
    text = paragraph.text.strip()
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": text,
    }


def table_record(table, index: int) -> dict:
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip().replace("\n", " / ") for cell in row.cells])
    return {"index": index, "rows": rows}


def extract(source: Path, output_dir: Path, slug: str) -> dict:
    document = Document(source)
    paragraphs = [
        paragraph_record(paragraph, index)
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if paragraph.text.strip()
    ]
    tables = [table_record(table, index) for index, table in enumerate(document.tables, start=1)]

    image_relations = []
    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            image_relations.append(
                {
                    "relationship_id": rel.rId,
                    "target": str(rel.target_ref),
                }
            )

    sections = []
    for index, section in enumerate(document.sections, start=1):
        sections.append(
            {
                "index": index,
                "width_pt": round(section.page_width.pt, 2),
                "height_pt": round(section.page_height.pt, 2),
                "top_margin_pt": round(section.top_margin.pt, 2),
                "bottom_margin_pt": round(section.bottom_margin.pt, 2),
                "left_margin_pt": round(section.left_margin.pt, 2),
                "right_margin_pt": round(section.right_margin.pt, 2),
            }
        )

    metadata = {
        "source_name": source.name,
        "size_bytes": source.stat().st_size,
        "sha256": sha256(source),
        "paragraph_count_nonempty": len(paragraphs),
        "table_count": len(tables),
        "image_count": len(image_relations),
        "section_count": len(sections),
        "sections": sections,
        "images": image_relations,
    }

    markdown = [f"# {source.name}", "", "## Paragraphs", ""]
    for item in paragraphs:
        markdown.append(
            f"P{item['index']:04d} [{item['style'] or 'No style'}] {item['text']}"
        )

    for table in tables:
        markdown.extend(["", f"## Table {table['index']}", ""])
        for row_index, row in enumerate(table["rows"], start=1):
            markdown.append(f"R{row_index:03d}: " + " | ".join(row))

    (output_dir / f"{slug}.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")
    (output_dir / f"{slug}.json").write_text(
        json.dumps(
            {"metadata": metadata, "paragraphs": paragraphs, "tables": tables},
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    return metadata


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output_dir", type=Path)
    parser.add_argument("inputs", nargs="+", help="slug=path")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    inventory = []
    for item in args.inputs:
        slug, raw_path = item.split("=", 1)
        inventory.append(extract(Path(raw_path), args.output_dir, slug))

    (args.output_dir / "docx-inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
