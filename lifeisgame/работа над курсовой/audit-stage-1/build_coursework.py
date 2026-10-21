from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
AUDIT = ROOT / "работа над курсовой" / "audit-stage-1"
SOURCE = ROOT / "работа над курсовой" / "Курсовая работа Глеба Короткова - исправленная.docx"
OUTPUT = ROOT / "работа над курсовой" / "Курсовая работа Глеба Короткова — редакция Luma.docx"
SAFE_DIR = Path("/Users/glebkorotkov/.codex/visualizations/2026/09/15/01a0a532-1dac-7093-b0b2-b15f17cf1a29/luma-coursework-audit")
SAFE_OUTPUT = SAFE_DIR / "Курсовая работа — редакция Luma.docx"


def clean_inline(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^>\s*", "", text)
    text = text.replace("**", "").replace("*", "").replace("`", "")
    text = re.sub(r"\[Автору:[^\]]+\]", "", text)
    text = re.sub(r"\[ИСТОЧНИК:[^\]]+\]", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def parse_markdown(path: Path):
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        if not raw.strip():
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", raw)
        if heading:
            blocks.append(("heading", len(heading.group(1)), clean_inline(heading.group(2))))
            i += 1
            continue
        if raw.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [clean_inline(cell) for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if re.match(r"^[-*]\s+", raw) or re.match(r"^\d+[.)]\s+", raw):
            items = []
            ordered = bool(re.match(r"^\d+[.)]\s+", raw))
            while i < len(lines):
                line = lines[i].strip()
                m = re.match(r"^(?:[-*]|\d+[.)])\s+(.+)$", line)
                if not m:
                    break
                items.append(clean_inline(m.group(1)))
                i += 1
            blocks.append(("list", ordered, items))
            continue
        para = [raw.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt.strip() or re.match(r"^#{1,3}\s+", nxt) or nxt.lstrip().startswith("|") or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+[.)]\s+", nxt):
                break
            para.append(nxt.strip())
            i += 1
        blocks.append(("paragraph", clean_inline(" ".join(para))))
    return blocks


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc: Document, rows: list[list[str]], caption: str | None = None):
    traceability = bool(rows and rows[0][:2] == ["ID", "Краткий предмет"])
    if traceability:
        caption = "Таблица 3.1 - Компактная матрица трассировки требований к реализации и проверке"
        short_evidence = {
            "FR-01": "Валидаторы, AuthService, SessionManager; сценарный прогон отсутствует.",
            "FR-02": "ProfileViewModel, Face ID, AvatarStore; редактирование учётных данных не реализовано.",
            "FR-03": "CalendarViewModel и TaskService; операции только с собственными задачами.",
            "FR-04": "EventKit-сервис чтения; разрешения и обновление требуют runtime-проверки.",
            "FR-05": "Формы и валидаторы; фактический конец короткого события не исправляется.",
            "FR-06": "TaskAutoScheduler; резервный путь не ограничен 23:00 и дедлайном.",
            "FR-07": "SubtaskAutoScheduler; атомарность общего сохранения не подтверждена.",
            "FR-08": "Диалог и splitTaskAtLateBoundary; связь по названию, сохранение нетранзакционно.",
            "FR-09": "LocalNotificationService; перепланирование после изменений не подтверждено.",
            "FR-10": "Одна из шести эмоций, шкала 1–10, сон через полночь.",
            "FR-11": "Таймер и Live Activity; внешнее событие ошибочно доступно для выбора.",
            "FR-12": "AnalysisService и UserAnalysisEngine; тестовые методы не запускались.",
            "FR-13": "weeklyStats и экраны; агрегаты и даты не проверены прогоном.",
            "FR-14": "AchievementRepository; переходы и границы уровней не проверены прогоном.",
            "NFR-01": "Проверки доступности есть; матрица устройств и версий ОС не пройдена.",
            "NFR-02": "Единые компоненты; полный аудит доступности не выполнен.",
            "NFR-03": "Таймер и обработка аватара; количественные замеры отсутствуют.",
            "NFR-04": "Миграция и rollback есть; ошибки сохранения частично подавляются.",
            "NFR-05": "Связи owner/parentTask есть; многопользовательские и DST-тесты отсутствуют.",
            "NFR-06": "Хеш, соль, Keychain, биометрия; отрицательные ветви не испытаны.",
            "NFR-07": "Разрешения запрашиваются по функции; динамическая проверка не выполнена.",
            "NFR-08": "Локальные данные и музыка; ресурсные прогоны отсутствуют.",
            "NFR-09": "MVVM-C, сервисы, репозитории и протоколы; часть служб - singleton.",
            "NFR-10": "Test bundle компилируется; Test action не настроен.",
        }
        rows = [rows[0]] + [[row[0], row[1], row[2], short_evidence.get(row[0], row[3])] for row in rows[1:]]
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(caption)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    try:
        table.style = "Table Grid"
    except KeyError:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:color"), "000000")
            borders.append(element)
        tbl_pr.append(borders)
    table.autofit = True
    for r_index, source_row in enumerate(rows):
        row = table.rows[r_index]
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
        prevent_row_split(row)
        if r_index == 0:
            set_repeat_table_header(row)
        for c_index in range(width):
            cell = row.cells[c_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = source_row[c_index] if c_index < len(source_row) else ""
            cell.text = text
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(8.5 if traceability else 9.5)
                    run.bold = r_index == 0
    doc.add_paragraph()


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for name, size, centered in [("Heading 1", 14, True), ("Heading 2", 13, False), ("Heading 3", 12, False)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True


def add_heading(doc: Document, text: str, level: int):
    if text == "2 Проектирование Luma":
        level = 1
    if text == "ЗАКЛЮЧЕНИЕ":
        level = 1
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body_paragraph(doc: Document, text: str):
    text = clean_inline(text)
    if not text:
        return None
    replacements = {
        "рисунок 3.12": "рисунок 3.4",
        "рисунке 3.12": "рисунке 3.4",
        "рисунок 3.20": "рисунок 3.5",
        "рисунке 3.20": "рисунке 3.5",
        "рисунок 3.22": "рисунок 3.6",
        "рисунке 3.22": "рисунке 3.6",
        "рисунках 3.12–3.22": "рисунках 3.4–3.6",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    if "Также в реализации нет эмоции «страх»" in text:
        text = "Проверка модели показала, что за одно сохранение пользователь выбирает одну эмоцию из шести: радость, грусть, злость, тревогу, покой или любовь. Интенсивность при сохранении ограничивается диапазоном 1–10, причина может отсутствовать. DiaryRepository связывает запись с пользователем и календарным днём. Для сна время относится к выбранному дню, а пробуждение переносится на следующие сутки, если оно не позже времени засыпания; после этого длительность вычисляется в минутах."
    if "Поэтому раннее описание" in text and "EmotionKind" in text:
        text = "Эмоциональная часть использует тот же EmotionKind, что и дневник. Интенсивность умножается на направление эмоции: радость, покой и любовь поддерживают ресурс, а грусть, злость и тревога увеличивают напряжение. Конкретные веса и пороги являются авторской эвристикой приложения, а не нормативными медицинскими показателями. Модуль формирует пользовательскую подсказку для самоорганизации и не выполняет диагностику состояния здоровья."
    p = doc.add_paragraph()
    p.add_run(text)
    return p


def add_blocks(doc: Document, blocks, skip_headings=(), stop_heading=None, start_heading=None, heading_level_override=None):
    active = start_heading is None
    pending_caption = None
    for block in blocks:
        if block[0] == "heading":
            _, level, title = block
            if title == start_heading:
                active = True
            if stop_heading and title == stop_heading:
                break
            if not active:
                continue
            if title in skip_headings:
                continue
            add_heading(doc, title, heading_level_override.get(title, level) if heading_level_override else level)
        elif not active:
            continue
        elif block[0] == "paragraph":
            text = block[1]
            if re.match(r"^(Таблица|Рисунок)\s+", text):
                if text.startswith("Таблица"):
                    pending_caption = text.replace("—", "-")
                continue
            if "Остальные диаграммы ответов следует перенести" in text:
                add_body_paragraph(doc, "В основном тексте приведены три диаграммы, отражающие понимание назначения приложения, удобство форм и календаря и общую оценку интерфейса. Первичная выгрузка Google Forms должна храниться вместе с материалами исследования для последующей проверки расчётов.")
                continue
            add_body_paragraph(doc, text)
        elif block[0] == "list":
            _, ordered, items = block
            for index, item in enumerate(items, 1):
                prefix = f"{index}. " if ordered else "- "
                add_body_paragraph(doc, prefix + item)
        elif block[0] == "table":
            add_table(doc, block[1], pending_caption)
            pending_caption = None


def add_figure(doc: Document, image_path: Path, caption: str, width_cm: float, keep_with_next=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = False
    r = cap.add_run(caption.replace("—", "-"))
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def add_figure_grid(doc: Document, items: list[tuple[Path, str]], width_cm: float = 5.0):
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)
    for index, (image_path, caption) in enumerate(items):
        row_index, column_index = divmod(index, 2)
        cell = table.cell(row_index, column_index)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(image_path), width=Cm(width_cm))
        cap = cell.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.first_line_indent = Cm(0)
        cap.paragraph_format.line_spacing = 1.0
        cap.paragraph_format.space_after = Pt(2)
        run = cap.add_run(caption)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)


def slice_blocks(blocks, start_title: str, end_title: str | None = None):
    result = []
    active = False
    for block in blocks:
        if block[0] == "heading":
            title = block[2]
            if title == start_title:
                active = True
            elif active and end_title and title == end_title:
                break
        if active:
            result.append(block)
    return result


def remove_body_from(doc: Document, paragraph):
    body = doc._element.body
    start = list(body).index(paragraph._p)
    for child in list(body)[start:]:
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def set_update_fields(doc: Document):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    doc = Document(SOURCE)
    original_references = [p.text.strip() for p in doc.paragraphs[355:379] if p.text.strip()]
    intro_p = next(p for p in doc.paragraphs if p.text.strip() == "ВВЕДЕНИЕ")
    remove_body_from(doc, intro_p)
    configure_styles(doc)
    set_update_fields(doc)

    intro = parse_markdown(AUDIT / "humanized-intro-ch1.md")
    competitor = parse_markdown(AUDIT / "competitor-sources.md")
    ch2 = parse_markdown(AUDIT / "humanized-ch2.md")
    ch3 = parse_markdown(AUDIT / "humanized-ch3.md")
    survey = parse_markdown(AUDIT / "humanized-survey-conclusion.md")

    add_blocks(doc, slice_blocks(intro, "ВВЕДЕНИЕ", "1.2 Анализ приложений-аналогов"))

    add_heading(doc, "1.2 Анализ приложений-аналогов", 2)
    comp_section = slice_blocks(competitor, "Рекомендуемый текст раздела 1.2", "Источники для добавления в список литературы")
    for block in comp_section:
        if block[0] == "heading":
            continue
        if block[0] == "paragraph" and block[1].startswith("Таблица 1.1"):
            continue
        if block[0] == "table":
            add_table(doc, block[1], "Таблица 1.1 - Сравнение функциональности приложений-аналогов")
        elif block[0] == "paragraph":
            add_body_paragraph(doc, block[1])

    sec_13_15 = slice_blocks(intro, "1.3 Функциональные требования", "1.6 Модель предметной области")
    add_blocks(doc, sec_13_15)

    uml_titles = [
        "Регистрация и авторизация",
        "Профиль пользователя",
        "Календарь и управление задачами",
        "Создание задач и автоматическое планирование",
        "Общая и подробная статистика",
        "Настройка и получение уведомлений",
        "Дневник эмоций и сна",
        "Фокус-сессия и Live Activity",
        "Анализ состояния пользователя",
        "Достижения, опыт и уровни",
    ]
    uml_files = sorted((AUDIT / "uml-rebuilt").glob("[0-9][0-9]-*.png"))
    for index, (image_path, title) in enumerate(zip(uml_files, uml_titles), 1):
        add_figure(doc, image_path, f"Рисунок 1.{index} - {title}", 10.8)

    add_blocks(doc, slice_blocks(intro, "1.6 Модель предметной области"))

    add_blocks(doc, ch2, heading_level_override={"2 Проектирование Luma": 1})

    ch3_sections = [
        ("3 Реализация и проверка приложения Luma", "3.1 Учётная запись, профиль и сессия"),
        ("3.1 Учётная запись, профиль и сессия", "3.2 Календарь и управление элементами"),
        ("3.2 Календарь и управление элементами", "3.3 Автоматическое планирование и граница 23:00"),
        ("3.3 Автоматическое планирование и граница 23:00", "3.4 Уведомления и дневник"),
        ("3.4 Уведомления и дневник", "3.5 Фокус-сессия и Live Activity"),
        ("3.5 Фокус-сессия и Live Activity", "3.6 Анализ дня, статистика и достижения"),
        ("3.6 Анализ дня, статистика и достижения", "3.7 Трассировка требований и результаты проверки"),
        ("3.7 Трассировка требований и результаты проверки", None),
    ]
    media = AUDIT / "template" / "media"
    for start, end in ch3_sections:
        add_blocks(doc, slice_blocks(ch3, start, end))
        if start == "3.2 Календарь и управление элементами":
            add_figure(doc, media / "image37.png", "Рисунок 3.1 - Календарь и собственные элементы пользователя", 8.1)
        elif start == "3.3 Автоматическое планирование и граница 23:00":
            add_figure(doc, media / "image2.jpg", "Рисунок 3.2 - Результат разделения задачи при выходе за границу дня", 8.1)
        elif start == "3.6 Анализ дня, статистика и достижения":
            add_figure(doc, media / "image4.jpg", "Рисунок 3.3 - Анализ дня и рекомендации", 8.1)

    survey_text = (AUDIT / "humanized-survey-conclusion.md").read_text(encoding="utf-8")
    survey_lines = survey_text.splitlines()
    for index, line in enumerate(survey_lines):
        if line.startswith("Сбор ответов выполнялся"):
            survey_lines[index] = "Сбор ответов выполнялся с помощью Google Forms. В исследовании приняли участие 14 респондентов. Участники оценивали готовое приложение Luma. Период проведения, единый набор пользовательских действий, продолжительность знакомства и состав устройств отдельно не фиксировались."
    survey_text = "\n".join(survey_lines)
    survey_text = re.sub(
        r"\[Автору: если при проведении фиксировались.*?как ограничение\.\]",
        "Способ набора участников и демографические характеристики не фиксировались, что ограничивает переносимость выводов.",
        survey_text,
        flags=re.S,
    )
    survey_text = survey_text.replace("## 3.4 Пользовательское исследование", "## 3.8 Пользовательское исследование")
    survey_text = survey_text.replace("## ЗАКЛЮЧЕНИЕ", "# ЗАКЛЮЧЕНИЕ")
    survey_clean = AUDIT / "survey-clean.md"
    survey_clean.write_text(survey_text, encoding="utf-8")
    survey_blocks = parse_markdown(survey_clean)
    survey_section = slice_blocks(survey_blocks, "3.8 Пользовательское исследование восприятия интерфейса", "ЗАКЛЮЧЕНИЕ")
    add_blocks(doc, survey_section)
    add_figure(doc, media / "image10.png", "Рисунок 3.4 - Оценка понимания назначения приложения", 13.8)
    add_figure(doc, media / "image20.png", "Рисунок 3.5 - Оценка удобства форм, календаря, выбора дат и задач", 13.8)
    add_figure(doc, media / "image23.png", "Рисунок 3.6 - Общая оценка интерфейса приложения", 13.8)
    add_blocks(doc, slice_blocks(survey_blocks, "ЗАКЛЮЧЕНИЕ"))

    add_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ И ЛИТЕРАТУРЫ", 1)
    for reference in original_references:
        add_body_paragraph(doc, reference)
    competitor_refs = [
        "25. TickTick: A To-Do List and Calendar to Keep You Organized [Электронный ресурс] // TickTick. - URL: https://www.ticktick.com/ (дата обращения: 25.09.2026). - Текст : электронный.",
        "26. Structured: Daily Planner [Электронный ресурс] // Structured. - URL: https://structured.app/ (дата обращения: 25.09.2026). - Текст : электронный.",
        "27. How to Create, Edit & Delete Tasks [Электронный ресурс] // Structured Help Center. - URL: https://help.structured.app/en/articles/338050 (дата обращения: 25.09.2026). - Текст : электронный.",
        "28. Simplify Task Management with Todoist [Электронный ресурс] // Todoist. - URL: https://www.todoist.com/task-management (дата обращения: 25.09.2026). - Текст : электронный.",
        "29. The Fabulous: Build Better Habits & Achieve Your Goals [Электронный ресурс] // Fabulous. - URL: https://www.thefabulous.co/ (дата обращения: 25.09.2026). - Текст : электронный.",
        "30. How Does Fabulous Work? [Электронный ресурс] // The Fabulous Help Center. - URL: https://help.thefabulous.co/en/support/solutions/articles/101000427430-how-does-fabulous-work- (дата обращения: 25.09.2026). - Текст : электронный.",
    ]
    for reference in competitor_refs:
        add_body_paragraph(doc, reference)

    add_heading(doc, "ПРИЛОЖЕНИЕ А", 1)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.first_line_indent = Cm(0)
    r = sub.add_run("Дополнительные экраны приложения")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    appendix = [
        ("image1.jpg", "Рисунок А.1 - История уведомлений"),
        ("image26.jpg", "Рисунок А.2 - Статистика эмоционального состояния"),
        ("image29.jpg", "Рисунок А.3 - Статистика выполнения задач"),
        ("image32.jpg", "Рисунок А.4 - Статистика сна"),
    ]
    add_figure_grid(doc, [(media / filename, caption) for filename, caption in appendix])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    SAFE_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUTPUT, SAFE_OUTPUT)
    print(OUTPUT)
    print(SAFE_OUTPUT)


if __name__ == "__main__":
    main()
