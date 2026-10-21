from __future__ import annotations

import copy
import os
import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("/Users/glebkorotkov/github/lifechange/lifeisgame/работа над курсовой")
SOURCE = ROOT / "Курсовая работа Глеба Короткова — правки UI и форматирование.docx"
OUTPUT = ROOT / "Курсовая работа Глеба Короткова — исправления по комментариям.docx"


def iter_all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def find_exact(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_prefix(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found by prefix: {prefix}")


def clear_runs(paragraph: Paragraph) -> None:
    # TOC paragraphs may keep their visible text inside hyperlinks/fields rather
    # than direct runs. Preserve only paragraph properties and rebuild content.
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_text(paragraph: Paragraph, text: str, bold_prefix: str | None = None) -> None:
    clear_runs(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def insert_before(target: Paragraph, text: str, style: str = "Normal", bold_prefix: str | None = None) -> Paragraph:
    paragraph = target.insert_paragraph_before(style=style)
    set_text(paragraph, text, bold_prefix=bold_prefix)
    return paragraph


def insert_after(anchor: Paragraph, text: str, style: str = "Normal", bold_prefix: str | None = None) -> Paragraph:
    p_elm = OxmlElement("w:p")
    anchor._p.addnext(p_elm)
    paragraph = Paragraph(p_elm, anchor._parent)
    paragraph.style = style
    set_text(paragraph, text, bold_prefix=bold_prefix)
    return paragraph


def delete_paragraph(paragraph: Paragraph) -> None:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)
    paragraph._p = paragraph._element = None


def rebuild_toc(doc: Document) -> None:
    entries = [
        ("toc 1", "Введение", 3),
        ("toc 1", "1 Анализ предметной области", 4),
        ("toc 2", "1.1 Особенности планирования задач с учётом состояния пользователя", 4),
        ("toc 2", "1.2 Анализ приложений-аналогов", 5),
        ("toc 2", "1.3 Нефункциональные требования", 7),
        ("toc 2", "1.4 Функциональные возможности приложения", 9),
        ("toc 3", "1.4.1 Текстовые сценарии ключевых вариантов использования", 16),
        ("toc 2", "1.5 Модель предметной области", 18),
        ("toc 1", "2 Проектирование", 19),
        ("toc 2", "2.1 Выбор платформы и средств разработки", 19),
        ("toc 3", "2.1.1 Выбор UI-фреймворка", 19),
        ("toc 3", "2.1.2 Выбор инструментов разработки", 20),
        ("toc 2", "2.2 Архитектурный подход", 20),
        ("toc 3", "2.2.1 Распределение ответственности", 21),
        ("toc 3", "2.2.2 Обоснование выбора MVVM-C", 21),
        ("toc 3", "2.2.3 Преимущества выбранного подхода", 22),
        ("toc 2", "2.3 Хранение данных и системные интеграции", 22),
        ("toc 2", "2.4 Проектирование пользовательского интерфейса", 24),
        ("toc 1", "3 Реализация приложения", 26),
        ("toc 2", "3.1 Учётная запись, профиль и сессия", 26),
        ("toc 2", "3.2 Календарь и управление элементами", 29),
        ("toc 2", "3.3 Автоматическое планирование и граница 23:00", 31),
        ("toc 2", "3.4 Уведомления и дневник", 35),
        ("toc 2", "3.5 Фокус-сессия и Live Activity", 36),
        ("toc 2", "3.6 Анализ дня, статистика и достижения", 37),
        ("toc 2", "3.7 Пользовательское исследование восприятия интерфейса", 43),
        ("toc 1", "Заключение", 47),
        ("toc 1", "Список использованной литературы", 48),
        ("toc 1", "Приложение А", 51),
        ("toc 1", "Приложение Б", 54),
    ]
    toc_paragraphs = [p for p in doc.paragraphs if p.style and p.style.name.lower().startswith("toc")]
    if len(toc_paragraphs) != len(entries):
        raise RuntimeError(f"Unexpected TOC size: {len(toc_paragraphs)} paragraphs, expected {len(entries)}")
    for paragraph, (style, title, page) in zip(toc_paragraphs, entries):
        paragraph.style = style
        set_text(paragraph, f"{title}\t{page}")
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
        fmt.keep_with_next = False
        fmt.tab_stops.clear_all()
        fmt.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.underline = False

    # Turn the visible contents into the cached result of a real Word TOC
    # field. Word can then rebuild it from Heading 1–3 automatically.
    first = toc_paragraphs[0]
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    begin_run.append(begin)
    begin_run.append(instruction)
    begin_run.append(separate)
    first._p.insert(1 if first._p.pPr is not None else 0, begin_run)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    toc_paragraphs[-1]._p.append(end_run)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def split_table_with_continuation(doc: Document, table_index: int, first_part_last_row: int, label: str) -> None:
    table = doc.tables[table_index]
    original = table._tbl

    def clone_with_rows(indices: list[int]):
        clone = copy.deepcopy(original)
        rows = list(clone.findall(qn("w:tr")))
        keep = set(indices)
        for index, row in enumerate(rows):
            if index not in keep:
                clone.remove(row)
        return clone

    row_count = len(table.rows)
    first = clone_with_rows(list(range(0, first_part_last_row + 1)))
    second = clone_with_rows([0] + list(range(first_part_last_row + 1, row_count)))

    # The supervisor requires the break to be shown by opening only the last
    # cell of the final row in the first fragment, not the whole row.
    last_row = first.findall(qn("w:tr"))[-1]
    last_cell = last_row.findall(qn("w:tc"))[-1]
    tc_pr = last_cell.find(qn("w:tcPr"))
    if tc_pr is None:
        tc_pr = OxmlElement("w:tcPr")
        last_cell.insert(0, tc_pr)
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "nil")

    label_elm = OxmlElement("w:p")
    label_p = Paragraph(label_elm, table._parent)
    label_p.style = "Normal"
    label_p.add_run(label)
    label_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    label_p.paragraph_format.first_line_indent = Cm(0)
    label_p.paragraph_format.keep_with_next = True
    label_p.paragraph_format.page_break_before = True

    original.addprevious(first)
    original.addprevious(label_elm)
    original.addprevious(second)
    original.getparent().remove(original)


def expand_citation(content: str) -> list[int]:
    content = content.strip()
    if ";" in content or "," in content:
        return [int(x) for x in re.split(r"\s*[;,]\s*", content) if x]
    match = re.fullmatch(r"(\d+)\s*[–-]\s*(\d+)", content)
    if match:
        start, end = map(int, match.groups())
        return list(range(start, end + 1))
    return [int(content)]


CITATION_RE = re.compile(r"\[(\d+(?:\s*(?:[;,]|[–-])\s*\d+)*)\]")


def rewrite_bibliography(doc: Document) -> None:
    bibliography_heading = find_exact(doc, "СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ")
    appendix_heading = find_exact(doc, "ПРИЛОЖЕНИЕ А")

    bibliography_paragraphs: list[Paragraph] = []
    started = False
    for paragraph in doc.paragraphs:
        if paragraph._p is bibliography_heading._p:
            started = True
            continue
        if paragraph._p is appendix_heading._p:
            break
        if started and re.match(r"^\d+\.\s+", paragraph.text.strip()):
            bibliography_paragraphs.append(paragraph)

    entries: dict[int, str] = {}
    for paragraph in bibliography_paragraphs:
        match = re.match(r"^(\d+)\.\s+(.*)$", paragraph.text.strip(), flags=re.S)
        if match:
            entries[int(match.group(1))] = match.group(2).strip()

    body_before_bibliography: list[Paragraph] = []
    for paragraph in doc.paragraphs:
        if paragraph._p is bibliography_heading._p:
            break
        body_before_bibliography.append(paragraph)

    order: list[int] = []
    for paragraph in body_before_bibliography:
        for match in CITATION_RE.finditer(paragraph.text):
            for number in expand_citation(match.group(1)):
                if number in entries and number not in order:
                    order.append(number)

    remap = {old: new for new, old in enumerate(order, 1)}

    def citation_sub(match: re.Match[str]) -> str:
        old_numbers = [n for n in expand_citation(match.group(1)) if n in remap]
        new_numbers = [remap[n] for n in old_numbers]
        if not new_numbers:
            return match.group(0)
        if len(new_numbers) > 2 and new_numbers == list(range(new_numbers[0], new_numbers[-1] + 1)):
            return f"[{new_numbers[0]}–{new_numbers[-1]}]"
        return "[" + "; ".join(str(n) for n in new_numbers) + "]"

    for paragraph in body_before_bibliography:
        if CITATION_RE.search(paragraph.text):
            set_text(paragraph, CITATION_RE.sub(citation_sub, paragraph.text))

    normalized_entries = []
    for new_number, old_number in enumerate(order, 1):
        text = entries[old_number]
        text = text.replace(" - ", " – ").replace(" — ", " – ")
        normalized_entries.append(f"{new_number}. {text}")

    if len(normalized_entries) > len(bibliography_paragraphs):
        raise RuntimeError("Not enough bibliography paragraphs")
    for paragraph, text in zip(bibliography_paragraphs, normalized_entries):
        set_text(paragraph, text)
    for paragraph in bibliography_paragraphs[len(normalized_entries):]:
        delete_paragraph(paragraph)


def set_font(run, name: str = "Times New Roman", size: int = 12) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def normalize_formatting(doc: Document) -> None:
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "toc 1", "toc 2", "toc 3"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if "Hyperlink" in doc.styles:
        doc.styles["Hyperlink"].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles["Hyperlink"].font.underline = False
    for style_name in ["Heading 2", "Heading 3"]:
        if style_name in doc.styles:
            doc.styles[style_name].font.bold = False

    for paragraph in iter_all_paragraphs(doc):
        for run in paragraph.runs:
            set_font(run)

        text = paragraph.text.strip()
        if paragraph.style and paragraph.style.name in {"Heading 2", "Heading 3"}:
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(1.25)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.space_after = Pt(12)
            for run in paragraph.runs:
                run.bold = False
        elif paragraph.style and paragraph.style.name == "Heading 1":
            paragraph.paragraph_format.space_after = Pt(24)
            if re.match(r"^\d+\s", text):
                paragraph.paragraph_format.left_indent = Cm(0)
                paragraph.paragraph_format.first_line_indent = Cm(1.25)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                paragraph.paragraph_format.left_indent = Cm(0)
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if text.startswith("Таблица ") or text.startswith("Продолжение таблицы ") or text.startswith("Окончание таблицы "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(0)
            paragraph.paragraph_format.first_line_indent = Cm(1.25)
        elif text.startswith("Рисунок "):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)

    # Remove trailing spaces before paragraph marks throughout the document.
    for paragraph in iter_all_paragraphs(doc):
        for run in reversed(paragraph.runs):
            if run.text:
                trimmed = run.text.rstrip()
                if trimmed != run.text:
                    run.text = trimmed
                break


def patch_alt_text(path: Path) -> None:
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    from lxml import etree

    with zipfile.ZipFile(path, "r") as source_zip:
        files = {name: source_zip.read(name) for name in source_zip.namelist()}

    root = etree.fromstring(files["word/document.xml"])
    paragraphs = root.xpath("//w:body/w:p | //w:body/w:tbl//w:p", namespaces=ns)
    texts = ["".join(p.xpath(".//w:t/text()", namespaces=ns)).strip() for p in paragraphs]

    for index, paragraph in enumerate(paragraphs):
        doc_properties = paragraph.xpath(".//wp:docPr", namespaces=ns)
        if not doc_properties:
            continue
        caption = ""
        for candidate in texts[index + 1:index + 12]:
            if candidate.startswith("Рисунок "):
                caption = candidate
                break
        if not caption:
            for candidate in reversed(texts[max(0, index - 11):index]):
                if candidate.startswith("Рисунок "):
                    caption = candidate
                    break
        description = caption or "Иллюстрация к курсовой работе Luma"
        for doc_property in doc_properties:
            doc_property.set("descr", description)

    # Hyperlinks and bibliography URLs must not retain Word's underline.
    for underline in root.xpath("//w:u", namespaces=ns):
        underline.set(f"{{{ns['w']}}}val", "none")

    # Remove orphan tab characters from otherwise empty paragraphs without
    # touching the tabs that align titles and page numbers in the contents.
    for tab in root.xpath("//w:p[not(.//w:t[normalize-space()])]//w:tab", namespaces=ns):
        tab.getparent().remove(tab)

    files["word/document.xml"] = etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone="yes",
    )

    ai_warning = "Контент, сгенерированный ИИ, может содержать ошибки".encode("utf-8")
    for name, data in list(files.items()):
        if name.endswith(".xml") and ai_warning in data:
            files[name] = data.replace(ai_warning, b"")

    fd, temporary_name = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(fd)
    try:
        with zipfile.ZipFile(temporary_name, "w", compression=zipfile.ZIP_DEFLATED) as target_zip:
            for name, data in files.items():
                target_zip.writestr(name, data)
        os.replace(temporary_name, path)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)


def scale_figure_before_caption(doc: Document, caption_prefix: str, factor: float) -> None:
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.startswith(caption_prefix) and index > 0:
            drawing_paragraph = paragraphs[index - 1]
            for extent in drawing_paragraph._p.xpath(".//wp:extent"):
                extent.set("cx", str(round(int(extent.get("cx")) * factor)))
                extent.set("cy", str(round(int(extent.get("cy")) * factor)))
            for extent in drawing_paragraph._p.xpath(".//a:xfrm/a:ext"):
                extent.set("cx", str(round(int(extent.get("cx")) * factor)))
                extent.set("cy", str(round(int(extent.get("cy")) * factor)))
            return
    raise RuntimeError(f"Figure caption not found: {caption_prefix}")


def main() -> None:
    doc = Document(SOURCE)
    try:
        set_text(find_exact(doc, "Глеб Коротков Дмитриевич"), "Коротков Глеб Дмитриевич")
    except ValueError:
        pass
    doc.core_properties.title = "Разработка iOS-приложения для адаптивного планирования задач с учётом состояния пользователя"
    doc.core_properties.subject = "Курсовая работа по проекту Luma"

    set_text(
        find_exact(doc, "Разработка iOS–приложения для адаптивного планирования и повышения продуктивности"),
        "Разработка iOS-приложения для адаптивного планирования задач с учётом состояния пользователя",
    )

    set_text(
        find_prefix(doc, "Работа выполняется в рамках заказной разработки для частного заказчика"),
        "Работа выполняется в рамках заказной разработки для частного заказчика, заинтересованного в создании персонального iOS-планировщика. Проектирование опирается на предоставленные технические данные о составе функций, правилах планирования и интеграции с системными средствами iOS. Результатом должна стать версия Luma, подготовленная для передачи заказчику.",
    )
    adaptive_anchor = find_prefix(doc, "Разрабатываемое приложение объединяет задачи")
    insert_after(
        adaptive_anchor,
        "Под адаптивным планированием в работе понимается формирование плана и рекомендаций с учётом календарной занятости, параметров задач, сведений о сне и эмоциональном состоянии пользователя.",
    )
    delete_paragraph(adaptive_anchor)
    set_text(
        find_prefix(doc, "Цель курсовой работы"),
        "Цель курсовой работы – разработать iOS-приложение Luma для адаптивного планирования задач с учётом состояния пользователя.",
    )
    set_text(
        find_prefix(doc, "5. реализовать и проверить"),
        "5. реализовать спроектированные функциональные возможности приложения.",
    )
    set_text(
        find_prefix(doc, "Практическая значимость работы"),
        "Практическая значимость работы состоит в создании Luma – мобильного приложения, которое объединяет задачи, календарные события, дневники состояния, фокус-сессии и статистику и может быть передано заказчику для практического использования.",
    )
    set_text(
        find_prefix(doc, "Исходной точкой проектирования Luma"),
        "Исходной точкой проектирования Luma служат технические данные и ограничения, предоставленные заказчиком. На этапе анализа они были сгруппированы по функциональным областям и сопоставлены с пользовательскими сценариями. Анализ приложений-аналогов дополняет постановку задачи, но не изменяет границы разрабатываемого приложения.",
    )

    comparison = find_prefix(doc, "Официальные описания не подтверждают")
    insert_after(comparison, "Результаты сравнения функциональности представлены в таблице 1.1.")

    set_text(
        find_prefix(doc, "Luma разрабатывается как нативное приложение"),
        "Luma разрабатывается как нативное приложение для iPhone. Приложение предназначено для устройств, поддерживающих используемые системные средства iOS, включая EventKit, LocalAuthentication, UserNotifications и ActivityKit. Возможности Face ID и Live Activity отображаются только на совместимых устройствах. Их отсутствие, а также отказ в доступе к системному календарю не должны блокировать работу с локальными задачами и независимыми разделами. Совместимость проверяется запуском приложения на заявленных конфигурациях и выполнением сценариев с недоступными или отозванными разрешениями.",
    )
    set_text(
        find_prefix(doc, "Таким образом, нефункциональные требования задают"),
        "Нефункциональные требования задают ограничения, которые учитываются при проектировании и реализации Luma. Их соблюдение оценивалось при ручной проверке основных пользовательских сценариев и анализе структуры приложения.",
    )

    set_text(
        find_prefix(doc, "Для фокус-сессии может быть выбран плейлист"),
        "Для фокус-сессии может быть выбран плейлист, содержащий один или несколько музыкальных треков. Музыка является дополнительной частью сценария и не влияет на возможность провести фокус-сессию без звукового сопровождения.",
    )
    set_text(
        find_prefix(doc, "Задача хранит сведения, необходимые для планирования"),
        "В предметной модели задача и событие рассматриваются как разные виды планируемого элемента. Задача имеет состояние выполнения, предполагаемую продолжительность и может переноситься или разделяться на подзадачи. Событие занимает фиксированный временной интервал и не участвует в автоматическом переносе; импортированное событие дополнительно хранит источник и доступно только для чтения. Подзадача связана с родительской задачей, а фокус-сессия – с тем действием, для которого пользователь проводил сеанс концентрации.",
    )

    scenarios_target = find_exact(doc, "1.5 Модель предметной области")
    insert_before(scenarios_target, "1.4.1 Текстовые сценарии ключевых вариантов использования", style="Heading 3")
    insert_before(
        scenarios_target,
        "Регистрация и вход. Предусловие: пользователь не авторизован. После открытия формы он вводит регистрационные данные либо электронную почту и пароль. Приложение проверяет обязательные поля, формат значений и наличие учётной записи. При корректных данных создаётся или восстанавливается сессия; при ошибке форма сохраняется, а пользователь получает сообщение о причине отказа. После первого входа можно включить Face ID, но неуспешная биометрическая проверка не исключает вход по паролю.",
        bold_prefix="Регистрация и вход.",
    )
    insert_before(
        scenarios_target,
        "Создание события. Предусловие: пользователь авторизован и открыл форму добавления элемента. Он указывает название и фиксированный интервал времени, при необходимости добавляет описание, важность и сложность. Приложение проверяет корректность интервала и сохраняет событие без автоматического переноса. Если обязательные данные отсутствуют или окончание указано раньше начала, сохранение не выполняется.",
        bold_prefix="Создание события.",
    )
    insert_before(
        scenarios_target,
        "Создание обычной задачи. Пользователь задаёт название, период выполнения, предполагаемую продолжительность, важность и сложность. Приложение получает занятые интервалы, формирует свободные промежутки и выбирает подходящее время без пересечений. Если выполнение выходит за 23:00, пользователь решает, оставить исходный интервал или перенести остаток на следующий день. Результатом становится задача с рассчитанным временем выполнения.",
        bold_prefix="Создание обычной задачи.",
    )
    insert_before(
        scenarios_target,
        "Создание сложной задачи. Пользователь задаёт общий период и добавляет не менее одной подзадачи с собственным названием, длительностью, важностью и сложностью. Приложение проверяет параметры и последовательно распределяет подзадачи по свободным интервалам. Если допустимого времени недостаточно, сохранение не завершается без сообщения пользователю. При успехе сохраняются родительская задача и связанные с ней подзадачи.",
        bold_prefix="Создание сложной задачи.",
    )
    insert_before(
        scenarios_target,
        "Учёт системного календаря. При первом обращении к календарным данным приложение запрашивает разрешение iOS. После согласия события выбранного периода загружаются как интервалы только для чтения и учитываются при отображении расписания и подборе времени задач. При отказе или отзыве разрешения приложение продолжает работать с локальными задачами и сообщает, что доступ можно изменить в настройках iOS.",
        bold_prefix="Учёт системного календаря.",
    )
    insert_before(
        scenarios_target,
        "Заполнение дневника. Авторизованный пользователь выбирает запись сна или эмоционального состояния. Для сна указываются время засыпания и пробуждения, после чего длительность рассчитывается автоматически, включая переход через полночь. Для эмоции выбираются вид и интенсивность от 1 до 10, а причина может быть пропущена. После сохранения запись связывается с текущим пользователем и календарным днём.",
        bold_prefix="Заполнение дневника.",
    )
    insert_before(
        scenarios_target,
        "Проведение фокус-сессии. Пользователь выбирает незавершённую задачу текущего дня, при желании включает локальный плейлист и запускает таймер. Приложение отображает оставшееся время и, если устройство поддерживает функцию, создаёт Live Activity. Завершение останавливает таймер и музыку, закрывает Live Activity, сохраняет фактическую длительность и отмечает задачу выполненной. Ошибка системной активности не должна останавливать обычный таймер.",
        bold_prefix="Проведение фокус-сессии.",
    )
    insert_before(
        scenarios_target,
        "Анализ дня. Пользователь открывает раздел анализа после добавления задач и, при наличии, дневниковых записей. Приложение получает задачи, события системного календаря, сведения о сне и эмоциях, рассчитывает индикативные показатели ресурса, нагрузки, баланса и полноты данных, затем формирует рекомендации. Если данных о состоянии недостаточно, результат помечается как предварительный. Анализ носит справочный характер и не используется для медицинской диагностики.",
        bold_prefix="Анализ дня.",
    )

    chapter2 = find_exact(doc, "2 Проектирование")
    insert_before(
        chapter2,
        "По итогам анализа определены границы Luma, ключевые пользовательские сценарии, требования к качеству и состав предметной модели. Эти результаты используются в следующей главе для выбора платформенных средств, архитектуры и способов хранения данных.",
    )

    set_text(find_exact(doc, "2.1.1 Выбор UI–фреймворка"), "2.1.1 Выбор UI-фреймворка")
    set_text(
        find_prefix(doc, "Техническое задание определяет Luma как нативное приложение"),
        "Технические данные определяют Luma как нативное приложение для iPhone. Поэтому выбор платформы не требовал сравнения с Android или кроссплатформенными решениями. Основные проектные решения касались UI-фреймворка, среды разработки, языка, хранения данных и системных API.",
    )
    set_text(
        find_prefix(doc, "SwiftUI выбран для экрана достижений"),
        "SwiftUI выбран для экрана достижений и интерфейса Live Activity, поскольку эти части преимущественно отображают состояние и требуют автоматического обновления представления при изменении данных. Декларативная модель SwiftUI позволяет описывать соответствие между состоянием и интерфейсом без ручного управления каждым обновлением элементов [34]. Для экрана достижений это упрощает отображение списка, прогресса и состояний блокировки, а для Live Activity соответствует модели представления ActivityKit и WidgetKit на экране блокировки и в Dynamic Island [41; 46]. Экран достижений встраивается в основной UIKit-интерфейс через UIHostingController. UIKit используется для сценариев со сложным управлением экраном и навигацией, а SwiftUI – для изолированных представлений, содержание которых определяется текущим состоянием данных.",
    )
    set_text(
        find_prefix(doc, "Использование системных средств уменьшает"),
        "Использование системных средств уменьшает число внешних зависимостей и обеспечивает согласованную работу с механизмами iOS. Перед обращением к защищённым данным приложение использует штатные запросы разрешений. Конкретная минимальная версия системы задаётся конфигурацией сборки и должна соответствовать доступности используемых API.",
    )

    chapter3 = find_exact(doc, "3 Реализация и проверка приложения")
    insert_before(
        chapter3,
        "В результате проектирования выбраны нативные средства iOS, сочетание UIKit и SwiftUI, архитектура MVVM-C, локальное хранение Core Data и отдельные сервисы системных интеграций. Такое разделение соответствует сценариям Luma и позволяет реализовать основную логику независимо от конкретных экранов.",
    )
    set_text(chapter3, "3 Реализация приложения")
    set_text(
        find_prefix(doc, "В результате разработки получено готовое мобильное приложение"),
        "В главе рассматривается реализация Luma через основные пользовательские сценарии. Для каждого сценария описаны входные данные, ключевой механизм обработки, результат и поведение при недоступности системной функции или некорректном вводе.",
    )
    set_text(
        find_prefix(doc, "Архитектура Luma включает представления UIKit"),
        "Общий поток обработки начинается с действия пользователя в UIKit- или SwiftUI-интерфейсе. Модель представления проверяет состояние экрана и передаёт данные прикладному сервису; сервис выполняет операцию через репозиторий либо системный API, а полученный результат возвращается в интерфейс. Координаторы отвечают за переходы, а DIContainer передаёт компонентам необходимые зависимости.",
    )
    set_text(
        find_prefix(doc, "Техническое задание предусматривает регистрацию"),
        "Сценарий учётной записи должен обеспечить регистрацию, обычный и биометрический вход, восстановление сессии и безопасный выход. Пользователь вводит имя, дату рождения, адрес электронной почты и пароль; приложение проверяет обязательные значения и создаёт локальную учётную запись. После входа профиль предоставляет доступ к данным пользователя, аватару, уровню, опыту и системным настройкам. На устройстве без Face ID биометрическая настройка скрывается, а ошибка проверки не блокирует вход по паролю. Основные экраны сценария показаны на рисунке 3.1.",
    )
    set_text(
        find_prefix(doc, "Пароль не записывается в модель Core Data"),
        "Пароль не записывается в модель Core Data. В текущей реализации KeychainManager вычисляет SHA-256-хеш из пароля и индивидуальной соли на основе UUID пользователя, а результат сохраняет в Keychain. Это исключает хранение пароля в открытом виде, однако SHA-256 является быстрой хеш-функцией общего назначения, а не специализированной функцией выработки ключа для паролей. Поэтому данное решение рассматривается как ограничение локального прототипа; перед эксплуатационным выпуском механизм следует заменить на PBKDF2, scrypt или Argon2 с документированными параметрами.",
    )
    set_text(
        find_prefix(doc, "Распределение ответственности построено следующим образом"),
        "При корректном вводе пользователь получает активную сессию и доступ к собственным данным. Ошибки формата, неизвестная учётная запись и неуспешная биометрическая проверка возвращаются на уровень формы без изменения сохранённых данных. После выхода сессия очищается, уведомления отменяются, а следующий доступ требует повторной авторизации.",
    )
    delete_paragraph(find_prefix(doc, "Регистрация, вход, проверка полей и восстановление сессии реализованы"))

    set_text(
        find_prefix(doc, "Техническое задание задаёт пять основных разделов"),
        "Календарь должен давать пользователю единое представление локальных задач и занятости из системного календаря. После выбора даты элементы сортируются по времени, фильтруются по состоянию выполнения и участвуют в расчёте дневного прогресса. Собственные задачи можно изменять, завершать и удалять; события iOS остаются доступными только для чтения. Последовательность взаимодействия с системным календарём приведена на рисунке А.2 приложения А.",
    )
    set_text(
        find_prefix(doc, "Техническое задание определяет три создаваемых вида элементов"),
        "При добавлении элемента пользователь выбирает событие, обычную или сложную задачу. Событие сохраняется в указанном интервале. Для обычной задачи приложение подбирает свободное время по длительности, важности и сложности, а сложную задачу предварительно раскладывает на подзадачи. Планирование учитывает локальные задачи и события системного календаря, исключает пересечения и преимущественно использует диапазон 08:00–23:00. Перенос остатка после 23:00 выполняется только с согласия пользователя. Алгоритм распределения обычной задачи представлен на рисунке 3.4.",
    )
    set_text(
        find_prefix(doc, "Техническое задание предусматривает включение уведомлений"),
        "Уведомления помогают не пропустить начало задачи и заполнение дневника, но используются только после разрешения пользователя. Для задач создаются напоминания за 30 и 5 минут, а для дневников – в заданные моменты дня; после сохранения записи оставшиеся напоминания этого типа отменяются. В дневнике эмоций сохраняются вид, интенсивность 1–10 и необязательная причина, а в дневнике сна – время засыпания и пробуждения с расчётом длительности, включая переход через полночь. Доставленные уведомления доступны в отдельном списке.",
    )
    set_text(
        find_prefix(doc, "Для фокус-сессии предусмотрены выбор незавершённой задачи"),
        "Фокус-сессия связывает обратный отсчёт с выбранной незавершённой задачей текущего дня. Пользователь может запустить сессию без музыки или выбрать локальный плейлист. Во время работы приложение обновляет таймер раз в секунду и, если функция поддерживается, передаёт состояние в Live Activity. Завершение останавливает таймер и музыку, сохраняет фактическую длительность и отмечает задачу выполненной; ошибка ActivityKit не должна прерывать основной сценарий.",
    )

    analysis_intro = find_prefix(doc, "Функция анализа дня предназначена")
    set_text(
        analysis_intro,
        "Анализ дня связывает план пользователя с доступными сведениями о его состоянии. При расчёте учитываются задачи текущего и следующего дня, сон, эмоции, просроченные элементы, уже выполненная работа и интервалы системного календаря. Модель является авторской эвристикой: она формирует ориентировочную оценку и рекомендации, не является медицинским инструментом и не подтверждает причинную связь между состоянием и результативностью. Подготовка показателей сна и эмоционального состояния показана на рисунке 3.8, последующие этапы – на рисунках Б.1–Б.4 приложения Б.",
    )
    insert_after(
        analysis_intro,
        "В расчёте используются четыре нормированных показателя. Ресурс – индикативная оценка доступного состояния по шкале от 0 до 100, полученная из записей сна и эмоций. Нагрузка – оценка объёма и трудоёмкости запланированной деятельности по шкале от 0 до 100 с учётом длительности, сложности, выполненной работы, просрочек и свободного времени. Баланс – приведённое к шкале 0–100 соотношение ресурса и нагрузки; значение 50 соответствует их условному равновесию. Достоверность – показатель полноты исходных данных по шкале от 0 до 100, а не статистическая точность или медицинская валидность результата.",
    )
    set_text(
        find_prefix(doc, "Ресурс дня рассчитывается"),
        "Ресурс рассчитывается по доступным сведениям о сне и эмоциональном состоянии. Оценка сна учитывает последнюю запись и среднее значение за семь дней, а рекомендуемая продолжительность выбирается по возрастной группе. В авторской эвристике сну присвоен вес 0,65, поскольку он представлен количественной продолжительностью, а эмоциональным данным – 0,35. Эти коэффициенты не являются клинически валидированными и рассматриваются как настраиваемые параметры прототипа. Если обе группы сведений отсутствуют, используется нейтральное значение 50.",
    )
    set_text(
        find_prefix(doc, "Нагрузка дня определяется"),
        "Нагрузка учитывает плановую продолжительность, сложность и важность задач, уже затраченное усилие, просрочки и фрагментацию списка. Интервалы системного календаря уменьшают доступное время. Итог ограничивается диапазоном 0–100. Баланс вычисляется как 50 + 0,65 × (ресурс − нагрузка) и также ограничивается шкалой 0–100; поэтому одинаковое число задач может давать разный результат при различной трудоёмкости и занятости дня.",
    )
    set_text(
        find_prefix(doc, "Достоверность результата зависит"),
        "Достоверность зависит только от полноты исходных данных. Начальное значение повышается при наличии даты рождения, записи сна за текущий день, не менее трёх записей сна за неделю и актуальной эмоциональной записи. Значения ниже 45 считаются низкими, от 45 до 74 – средними, от 75 – высокими. При низкой достоверности приложение показывает предварительный анализ и не подменяет отсутствующие наблюдения категоричным выводом.",
    )
    set_text(
        find_prefix(doc, "Открытие достижения влияет на систему опыта"),
        "Открытие достижения начисляет 50 единиц опыта, а каждые 250 единиц повышают уровень пользователя. Эти значения являются конфигурируемыми параметрами прототипа и не рассматриваются как доказанный способ повышения мотивации; их можно корректировать после анализа пользовательской обратной связи. На основании уровня выбирается звание и рассчитывается прогресс до следующего уровня, который отображается в достижениях и профиле.",
    )
    study_limit = find_prefix(doc, "При интерпретации результатов необходимо учитывать ограничения исследования")
    set_text(
        study_limit,
        "При интерпретации результатов необходимо учитывать ограничения исследования. Выборка включала 14 человек, поэтому оценки описывают мнение только данной группы. Способ набора участников, демографические характеристики, единый набор действий, продолжительность знакомства и состав устройств не фиксировались. До сверки с первичной выгрузкой Google Forms численные значения рассматриваются как авторская сводка. Поэтому результаты можно называть предварительной оценкой восприятия интерфейса, но не подтверждением полной готовности продукта.",
    )

    conclusion_heading = find_exact(doc, "ЗАКЛЮЧЕНИЕ")
    insert_before(
        conclusion_heading,
        "В главе показано, как основные пользовательские сценарии реализованы через интерфейс, модели представления, сервисы, репозитории и системные API iOS. Автор выполнил ручную проверку предусмотренных техническим заданием функций и ограничений. Пользовательское исследование при этом рассматривается отдельно как предварительная оценка восприятия интерфейса.",
    )
    set_text(
        find_prefix(doc, "В результате выполнения курсовой работы было разработано"),
        "В результате выполнения курсовой работы разработано iOS-приложение Luma для адаптивного планирования задач с учётом состояния пользователя. Работа выполнена для частного заказчика на основе предоставленных технических данных.",
    )
    set_text(
        find_prefix(doc, "Поставленное задание выполнено"),
        "Поставленное задание выполнено: предусмотренные техническим заданием функциональные и нефункциональные требования реализованы и проверены автором вручную. Пользовательское исследование дало предварительную оценку восприятия интерфейса, но не являлось приёмочным испытанием. Дальнейшая работа может включать замену прототипной схемы проверки пароля, калибровку коэффициентов анализа дня и уточнение интерфейса по результатам практического использования.",
    )

    bibliography_heading = find_exact(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ И ЛИТЕРАТУРЫ")
    set_text(bibliography_heading, "СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ")
    appendix_a = find_exact(doc, "ПРИЛОЖЕНИЕ А")
    widget_entry = insert_before(
        appendix_a,
        "46. WidgetKit [Электронный ресурс] // Apple Developer Documentation. – URL: https://developer.apple.com/documentation/widgetkit (дата обращения: 13.10.2026). – Текст : электронный.",
    )
    widget_entry.paragraph_format.first_line_indent = Cm(1.25)

    # Replace typographically incorrect dashes in compounds and remove em dashes.
    for paragraph in iter_all_paragraphs(doc):
        for run in paragraph.runs:
            replacement = run.text.replace("iOS–", "iOS-").replace("UI–", "UI-").replace("—", "–")
            # Reassigning an empty drawing run removes its embedded image in python-docx.
            if replacement != run.text:
                run.text = replacement

    # Add the new heading to the static table of contents; page number is updated after rendering.
    toc_anchor = find_prefix(doc, "1.4 Функциональные возможности приложения")
    toc_scenario = insert_after(toc_anchor, "1.4.1 Текстовые сценарии ключевых вариантов использования\t0", style="toc 3")
    toc_scenario.paragraph_format.first_line_indent = Cm(0)
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("toc"):
            if paragraph.text.startswith("3 Реализация и проверка приложения"):
                set_text(paragraph, paragraph.text.replace("3 Реализация и проверка приложения", "3 Реализация приложения"))
            if paragraph.text.startswith("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ И ЛИТЕРАТУРЫ"):
                set_text(paragraph, paragraph.text.replace("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ И ЛИТЕРАТУРЫ", "СПИСОК ИСПОЛЬЗОВАННОЙ ЛИТЕРАТУРЫ"))

    # Split the two long data tables so that continuation labels are explicit.
    # Keep the first fragment short enough to fit on one page; otherwise Word
    # leaves the last row of the fragment alone on the following page.
    split_table_with_continuation(doc, 13, 5, "Окончание таблицы 3.1")
    split_table_with_continuation(doc, 0, 5, "Окончание таблицы 1.1")

    rewrite_bibliography(doc)
    normalize_formatting(doc)
    for number in range(1, 6):
        paragraph = find_prefix(doc, f"{number}. ")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    rebuild_toc(doc)
    scale_figure_before_caption(doc, "Рисунок Б.1", 0.88)
    for paragraph in iter_all_paragraphs(doc):
        if paragraph.text and not paragraph.text.strip():
            for run in paragraph.runs:
                run.text = ""
    doc.save(OUTPUT)
    patch_alt_text(OUTPUT)

    with zipfile.ZipFile(OUTPUT, "r") as archive:
        ai_phrase = "Контент, сгенерированный ИИ, может содержать ошибки".encode("utf-8")
        if any(ai_phrase in archive.read(name) for name in archive.namelist() if name.endswith(".xml")):
            raise RuntimeError("AI warning metadata remains in the document")

    print(OUTPUT)


if __name__ == "__main__":
    main()
